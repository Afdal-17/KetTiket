import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook
wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styles
font_title = Font(name="Arial", size=14, bold=True, color="1F4E78")
font_subtitle = Font(name="Arial", size=11, italic=True, color="595959")
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Arial", size=10, bold=True)
font_regular = Font(name="Arial", size=10)
font_code = Font(name="Consolas", size=9)

fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fill_soft_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_light_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
fill_green_status = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fill_blue_status = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")

thin_border_side = Side(border_style="thin", color="D9D9D9")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

# ---------------------------------------------------------
# SHEET 1: ALUR NARASI
# ---------------------------------------------------------
ws1 = wb.create_sheet(title="ALUR NARASI")
ws1.views.sheetView[0].showGridLines = True

ws1.cell(row=1, column=1, value="DESKRIPSI DAN ALUR NARASI SISTEM KETTIKET (SaaS Concert Ticketing Platform)").font = font_title
ws1.cell(row=2, column=1, value="Sistem Tiketing Konser Musik berbasis SaaS dengan fitur E-Commerce, Seat Map, E-Ticket QR Code, serta Modul AI Intelligence Assistant.").font = font_subtitle

headers1 = ["Kategori", "Fitur / Komponen Utama", "Deskripsi Narasi & Alur Operasional Sistem"]
ws1.row_dimensions[4].height = 28
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=4, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center

narasi_data = [
    ("Gambaran Umum", "Platform SaaS Tiketing Konser", "Aplikasi KetTiket merupakan platform SaaS (Software as a Service) tiket konser musik yang memungkinkan Customer mencari, memilih tempat duduk (seat map), membeli tiket, dan mendapatkan E-Ticket berbasis QR Code. Penyelenggara (Organizer) dapat mengelola event, tipe tiket, denah kursi, dan pemindaian tiket, sedangkan Super Admin mengelola verifikasi organizer, venue, dan pemantauan platform."),
    ("Customer Workflow", "E-Commerce & Checkout Flow", "Customer mendaftar/login -> Melakukan pencarian & browsing event konser -> Melihat detail event, tanggal, dan lokasi venue -> Memilih kategori tiket & tempat duduk (seat map) -> Melakukan checkout order -> Pembayaran via Payment Gateway (Tripay / Midtrans) -> Sistem menerbitkan Order & E-Ticket lengkap dengan Barcode/QR Code -> E-Ticket dapat dilihat & diunduh di aplikasi."),
    ("Customer Workflow", "Check-in & Verification", "Saat hari H konser, Customer menunjukkan E-Ticket QR Code kepada petugas/scanner di venue -> Scanner membaca QR Code via endpoint API -> Sistem memvalidasi keaslian & status tiket (mencatat scan_log) -> Mencegah penggunaan tiket ganda (double entry)."),
    ("Organizer Workflow", "Event & Seat Management", "Organizer terverifikasi dapat membuat event baru -> Memilih/mengatur Venue dan Layout Seat Map (seksi, baris, nomor kursi) -> Mengatur Tipe Tiket & Kuota Harga -> Memantau Laporan Penjualan (Sales Report) secara real-time -> Melakukan Scan E-Ticket di lokasi acara."),
    ("Super Admin Workflow", "Platform & Tenant Management", "Super Admin mengelola pendaftaran & verifikasi data Organizer -> Mengelola master data User & Role -> Mengelola master data Venue & Kapasitas -> Mengatur Integrasi Payment Gateway -> Memantau Dashboard analytics transaksi seluruh platform."),
    ("AI Module", "Customer AI Assistant", "Fitur Chat AI untuk Customer yang memberikan rekomendasi konser musik berdasarkan preferensi, menjawab FAQ seputar acara, jadwal, lokasi venue, serta syarat & ketentuan pembelian tiket secara interaktif."),
    ("AI Module", "Organizer AI Assistant", "Fitur AI bagi Organizer untuk melakukan auto-generate deskripsi event yang menarik, menganalisis tren penjualan tiket, memprediksi permintaan penonton, serta memberikan rekomendasi strategi harga tiket (dynamic pricing)."),
    ("AI Module", "Super Admin AI Analytics", "Fitur AI bagi Super Admin untuk menganalisis performa platform secara keseluruhan, mendeteksi aktivitas mencurigakan/fraud transaksi pembelian tiket, serta mengukur kinerja masing-masing organizer.")
]

row_idx = 5
for cat, sub, desc in narasi_data:
    ws1.cell(row=row_idx, column=1, value=cat).font = font_bold
    ws1.cell(row=row_idx, column=2, value=sub).font = font_bold
    ws1.cell(row=row_idx, column=3, value=desc).font = font_regular
    
    ws1.cell(row=row_idx, column=1).alignment = align_top_left
    ws1.cell(row=row_idx, column=2).alignment = align_top_left
    ws1.cell(row=row_idx, column=3).alignment = align_top_left
    
    for c in range(1, 4):
        ws1.cell(row=row_idx, column=c).border = thin_border
    ws1.row_dimensions[row_idx].height = 50
    row_idx += 1

# ---------------------------------------------------------
# SHEET 2: FITUR & ENTITAS
# ---------------------------------------------------------
ws2 = wb.create_sheet(title="FITUR & ENTITAS")
ws2.views.sheetView[0].showGridLines = True

ws2.cell(row=1, column=1, value="DAFTAR FITUR UTAMA SISTEM").font = font_title

headers2_1 = ["No", "Modul / Fitur", "Deskripsi Fitur", "Pengguna / Actor"]
ws2.row_dimensions[3].height = 26
for c_idx, h in enumerate(headers2_1, 1):
    cell = ws2.cell(row=3, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center

fitur_list = [
    (1, "E-Commerce & Ticketing", "Browsing event konser, filter venue, pilih seat map/kursi, checkout order, dan pembayaran online.", "Customer"),
    (2, "E-Ticket & QR Check-in", "Penerbitan E-Ticket unik berpola Barcode/QR Code, serta scanner validasi tiket masuk di lokasi konser.", "Customer, Scanner / Organizer"),
    (3, "Event & Seat Map Management", "Manajemen pembuatan event konser, pengelompokan seksi/baris kursi venue, serta penentuan kuota & harga tiket.", "Organizer"),
    (4, "Organizer Verification & Admin Control", "Verifikasi legalitas organizer, pengelolaan pengguna (users/roles), master data venue, serta dashboard pemantauan.", "Super Admin"),
    (5, "Customer AI Assistant", "Asisten percakapan AI untuk rekomendasi konser, konsultasi event, dan informasi FAQ konser secara otomatis.", "Customer"),
    (6, "Organizer AI Business Intelligence", "Modul AI untuk pembuat deskripsi event otomatis, analisis tren penjualan tiket, dan rekomendasi penetapan harga.", "Organizer"),
    (7, "Super Admin AI Analytics & Anti-Fraud", "Pemantauan performa platform, analisis perilaku transaksi, dan deteksi dini aktivitas pembelian tiket yang mencurigakan.", "Super Admin")
]

row_idx = 4
for no, feat, desc, actor in fitur_list:
    ws2.cell(row=row_idx, column=1, value=no).alignment = align_center
    ws2.cell(row=row_idx, column=2, value=feat).font = font_bold
    ws2.cell(row=row_idx, column=3, value=desc).font = font_regular
    ws2.cell(row=row_idx, column=4, value=actor).alignment = align_center
    
    for c in range(1, 5):
        ws2.cell(row=row_idx, column=c).border = thin_border
    ws2.row_dimensions[row_idx].height = 32
    row_idx += 1

row_idx += 2
ws2.cell(row=row_idx, column=1, value="DAFTAR ENTITAS & AKTOR SISTEM").font = font_title
row_idx += 1

headers2_2 = ["Aktor / Role", "Tanggung Jawab & Hak Akses Utama dalam Sistem"]
ws2.row_dimensions[row_idx].height = 26
for c_idx, h in enumerate(headers2_2, 1):
    cell = ws2.cell(row=row_idx, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center
row_idx += 1

entitas_list = [
    ("Customer", "Mendaftar akun, mencari & melihat detail konser, memilih tempat duduk, melakukan pembayaran, menerima E-ticket QR Code, serta menggunakan AI Assistant untuk rekomendasi."),
    ("Organizer", "Mengelola profil perusahaan organizer, membuat & mengedit event konser, mengoperasikan seat map & tipe tiket, melihat laporan penjualan, mengoperasikan scan tiket, serta memanfaatkan AI Business Intelligence."),
    ("Super Admin", "Verifikasi dan persetujuan akun Organizer, mengelola data seluruh pengguna platform, mengelola master data Venue & Kapasitas, mengatur konfigurasi Payment Gateway, serta memantau analitik AI platform."),
    ("Payment Gateway", "Sistem pihak ketiga (Tripay / Midtrans) yang memproses transaksi pembayaran tiket dan mengirimkan status callback/webhook secara real-time."),
    ("Barcode / Ticket Scanner", "Aplikasi / perangkat pemindai yang memvalidasi QR Code E-Ticket pengunjung di pintu masuk konser via API Backend.")
]

for actor, resp in entitas_list:
    ws2.cell(row=row_idx, column=1, value=actor).font = font_bold
    ws2.cell(row=row_idx, column=1).alignment = align_center
    ws2.cell(row=row_idx, column=2, value=resp).font = font_regular
    ws2.cell(row=row_idx, column=2).alignment = align_left
    
    ws2.cell(row=row_idx, column=1).border = thin_border
    ws2.cell(row=row_idx, column=2).border = thin_border
    ws2.row_dimensions[row_idx].height = 30
    row_idx += 1


# ---------------------------------------------------------
# SHEET 3: THIRD PARTY & DIAGRAM
# ---------------------------------------------------------
ws3 = wb.create_sheet(title="THIRD PARTY & DIAGRAM")
ws3.views.sheetView[0].showGridLines = True

ws3.cell(row=1, column=1, value="INTEGRASI LAYANAN PIHAK KETIGA (THIRD PARTY SERVICES)").font = font_title

headers3_1 = ["Layanan Third Party", "Komponen Integrasi", "Keterangan / Endpoints / Dokumentasi"]
ws3.row_dimensions[3].height = 26
for c_idx, h in enumerate(headers3_1, 1):
    cell = ws3.cell(row=3, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center

third_party_data = [
    ("Tripay / Payment Gateway", "Base URL & Callback Webhook", "https://tripay.co.id/developer - Digunakan untuk simulasi & transaksi pembayaran tiket konser (QRIS, VA, E-Wallet)."),
    ("Barcode / QR Code Engine", "Generation & Verification API", "Menggunakan library QR Code generator (seperti Endroid/QrCode) untuk menerbitkan QR Code unik E-Ticket."),
    ("AI Engine (OpenAI / LLM)", "API Model Key & Prompt Handler", "Digunakan pada modul Customer AI Assistant, Organizer AI Generator, dan Super Admin AI Analytics."),
    ("Email / Notification", "SMTP / Mail Driver", "Pemberitahuan bukti konfirmasi order dan pengiriman berkas E-Ticket PDF ke email Customer.")
]

row_idx = 4
for tp, comp, desc in third_party_data:
    ws3.cell(row=row_idx, column=1, value=tp).font = font_bold
    ws3.cell(row=row_idx, column=2, value=comp).font = font_regular
    ws3.cell(row=row_idx, column=3, value=desc).font = font_regular
    
    ws3.cell(row=row_idx, column=1).alignment = align_left
    ws3.cell(row=row_idx, column=2).alignment = align_left
    ws3.cell(row=row_idx, column=3).alignment = align_left
    
    for c in range(1, 4):
        ws3.cell(row=row_idx, column=c).border = thin_border
    ws3.row_dimensions[row_idx].height = 28
    row_idx += 1

row_idx += 2
ws3.cell(row=row_idx, column=1, value="DOKUMEN PERANCANGAN DIAGRAM (UML & ERD)").font = font_title
row_idx += 1

headers3_2 = ["Nama Diagram", "File Berkas Perancangan (.puml)", "Deskripsi Cakupan Diagram"]
ws3.row_dimensions[row_idx].height = 26
for c_idx, h in enumerate(headers3_2, 1):
    cell = ws3.cell(row=row_idx, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center
row_idx += 1

diagram_data = [
    ("Core ERD Diagram", "perancangan UML/ERD.puml", "Mencakup 11 entitas utama: users, organizers, events, venues, seats, ticket_types, orders, order_details, payments, tickets, scan_logs."),
    ("AI Module ERD Diagram", "Perancangan AI/ERD ai modul.puml", "Mencakup entitas ai_conversations, ai_messages, ai_feedback, ai_logs dan relasinya dengan core system."),
    ("Core Use Case Diagram", "perancangan UML/usecase.puml", "Interaksi aktor Customer, Organizer, Super Admin, Payment Gateway, dan Scanner."),
    ("AI Use Case Diagram", "Perancangan AI/usecase ai.puml", "Interaksi aktor modul AI untuk Customer AI, Organizer AI, dan Super Admin AI."),
    ("Sequence & Activity Diagrams", "perancangan UML/ & Perancangan AI/", "Detail alur proses payment, check-in barcode, organizer event creation, serta interaksi percakapan AI.")
]

for diag, path, desc in diagram_data:
    ws3.cell(row=row_idx, column=1, value=diag).font = font_bold
    ws3.cell(row=row_idx, column=2, value=path).font = font_code
    ws3.cell(row=row_idx, column=3, value=desc).font = font_regular
    
    ws3.cell(row=row_idx, column=1).alignment = align_left
    ws3.cell(row=row_idx, column=2).alignment = align_left
    ws3.cell(row=row_idx, column=3).alignment = align_left
    
    for c in range(1, 4):
        ws3.cell(row=row_idx, column=c).border = thin_border
    ws3.row_dimensions[row_idx].height = 28
    row_idx += 1


# ---------------------------------------------------------
# SHEET 4: BREAKDOWN BACK-END TASKS
# ---------------------------------------------------------
ws4 = wb.create_sheet(title="Breakdown Back-End")
ws4.views.sheetView[0].showGridLines = True

ws4.cell(row=1, column=1, value="DAFTAR TASKS BACK-END DEVELOPER (KetTiket Platform)").font = font_title

headers4 = ["No", "Task Implementation", "Criteria & Requirements", "Status", "Scope / Module", "PIC / Owner", "Catatan / Controller File"]
ws4.row_dimensions[3].height = 28
for c_idx, h in enumerate(headers4, 1):
    cell = ws4.cell(row=3, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center

be_tasks = [
    (1, "Analisis & Perancangan Sistem KetTiket", "Analisis struktur ERD Core, ERD AI Module, Use Case Diagram, dan Arsitektur SaaS Tiketing Konser.", "Selesai", "System Architecture", "System Analyst", "Dokumen PlantUML & Database Schema"),
    (2, "Database Migration & Database Seeder", "Pembuatan migration tabel core (users, organizers, venues, seats, events, ticket_types, orders, order_details, payments, tickets, scan_logs) dan tabel AI (ai_conversations, ai_messages, ai_feedback, ai_logs).", "Selesai", "Database Core & AI", "Backend Dev", "2026_08_04_* & 2026_08_11_* migrations"),
    (3, "Model Declaration & Relationships", "Deklarasi Eloquent Model, mass assignment ($fillable), serta relasi antar entitas (belongsTo, hasMany, hasOne) sesuai ERD.", "Selesai", "Database Models", "Backend Dev", "App/Models/* (15 Models)"),
    (4, "Authentication & Authorization API", "Fitur register user/organizer, login JWT/Sanctum, social login, profil pengguna, dan pembatasan middleware auth:sanctum.", "Selesai", "Auth & Security", "Backend Dev", "AuthController.php"),
    (5, "Organizer Management & Verifikasi Admin", "Admin dapat melihat daftar pengajuan organizer, memverifikasi status organizer (verify), melihat detail saldo & transaksi organizer.", "Selesai", "Admin Panel", "Backend Dev", "AdminController.php"),
    (6, "Venue & Seat Map Management API", "Manajemen data venue (lokasi, kapasitas) dan auto-generate seat map (seksi, baris, nomor kursi).", "Selesai", "Organizer & Admin", "Backend Dev", "VenueController.php"),
    (7, "Event & Ticket Type Management API", "CRUD Event Konser (kategori, banner image, tanggal) dan pendaftaran tipe tiket (harga, kuota) oleh Organizer.", "Selesai", "Organizer Panel", "Backend Dev", "EventController.php"),
    (8, "Booking & Order Ticket Processing API", "Proses pemilihan tempat duduk, checkout pesanan tiket, perhitungan total pembayaran, dan penguncian kuota tiket.", "Selesai", "Customer Core", "Backend Dev", "OrderController.php"),
    (9, "Payment Gateway Integration & Webhook", "Simulasi & integrasi pembayaran tiket, callback webhook status pembayaran (Pending -> Paid -> Failed).", "Selesai", "Payment Module", "Backend Dev", "PaymentController.php"),
    (10, "E-Ticket Generation & QR Check-in API", "Penerbitan E-Ticket unik berpola Barcode/QR Code setelah pembayaran sukses, serta endpoint scan tiket untuk pemindaian masuk konser.", "Selesai", "Ticketing & Scanner", "Backend Dev", "TicketController.php"),
    (11, "Super Admin Dashboard & Reports API", "Endpoint dashboard analytics statistik penjualan, total transaksi, monitoring users/organizers, serta laporan keuangan.", "Selesai", "Admin Analytics", "Backend Dev", "AdminController.php"),
    (12, "Customer AI Assistant API", "Endpoint percakapan AI interaktif, rekomendasi konser otomatis berdasarkan preferensi, riwayat percakapan, dan feedback rating AI.", "Selesai", "AI Module", "Backend Dev", "CustomerAiController.php"),
    (13, "Organizer AI Business Intelligence API", "Endpoint AI untuk auto-generate deskripsi event, analisis tren penjualan tiket, rekomendasi kuota/harga tiket, dan pembuatan ringkasan laporan.", "Selesai", "AI Module", "Backend Dev", "OrganizerAiController.php"),
    (14, "Super Admin AI Analytics & Anti-Fraud API", "Endpoint AI untuk analisis performa platform secara menyeluruh dan deteksi dini aktivitas pesanan tiket yang mencurigakan.", "Selesai", "AI Module", "Backend Dev", "AdminAiController.php")
]

row_idx = 4
for no, task, crit, status, scope, pic, note in be_tasks:
    ws4.cell(row=row_idx, column=1, value=no).alignment = align_center
    ws4.cell(row=row_idx, column=2, value=task).font = font_bold
    ws4.cell(row=row_idx, column=3, value=crit).font = font_regular
    
    cell_status = ws4.cell(row=row_idx, column=4, value=status)
    cell_status.alignment = align_center
    cell_status.font = font_bold
    cell_status.fill = fill_green_status if status == "Selesai" else fill_blue_status
    
    ws4.cell(row=row_idx, column=5, value=scope).alignment = align_center
    ws4.cell(row=row_idx, column=6, value=pic).alignment = align_center
    ws4.cell(row=row_idx, column=7, value=note).font = font_code
    
    ws4.cell(row=row_idx, column=2).alignment = align_left
    ws4.cell(row=row_idx, column=3).alignment = align_left
    ws4.cell(row=row_idx, column=7).alignment = align_left
    
    for c in range(1, 8):
        ws4.cell(row=row_idx, column=c).border = thin_border
    ws4.row_dimensions[row_idx].height = 36
    row_idx += 1


# ---------------------------------------------------------
# SHEET 5: BREAKDOWN FRONT-END TASKS
# ---------------------------------------------------------
ws5 = wb.create_sheet(title="Breakdown Front-End")
ws5.views.sheetView[0].showGridLines = True

ws5.cell(row=1, column=1, value="DAFTAR TASKS FRONT-END / MOBILE DEVELOPER (KetTiket Application)").font = font_title

headers5 = ["No", "Modul / Halaman Screen", "Keterangan Implementation", "Status", "PIC / Owner"]
ws5.row_dimensions[3].height = 28
for c_idx, h in enumerate(headers5, 1):
    cell = ws5.cell(row=3, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center

fe_tasks = [
    (1, "Setup Project & UI Template", "Inisialisasi framework Front-End / Vue / Mobile App, konfigurasi router & state management.", "Selesai", "FrontEnd Dev"),
    (2, "Halaman Register & Login Customer", "Form registrasi akun, login email/password, social login, serta persetujuan syarat & ketentuan.", "Selesai", "FrontEnd Dev"),
    (3, "Dashboard Customer & Search Event", "Tampilan rekomendasi konser utama, pencarian event, filter lokasi venue, dan kategori musik.", "Selesai", "FrontEnd Dev"),
    (4, "Halaman Detail Event & Seat Map Picker", "Tampilan detail konser, deskripsi, tanggal, serta visual peta tempat duduk (seat map) interaktif.", "Selesai", "FrontEnd Dev"),
    (5, "Checkout & Payment Screen", "Ringkasan pesanan tiket, rincian biaya, pilihan metode pembayaran Tripay/QRIS, serta timer pembayaran.", "Selesai", "FrontEnd Dev"),
    (6, "Halaman E-Ticket & Download PDF", "Daftar tiket yang dibeli, tampilan QR Code E-Ticket unik, serta tombol unduh tiket.", "Selesai", "FrontEnd Dev"),
    (7, "Halaman Customer AI Assistant Chat", "Antarmuka percakapan interaktif AI Assistant untuk konsultasi event dan rekomendasi konser.", "Selesai", "FrontEnd Dev"),
    (8, "Dashboard Organizer & Event Creation", "Form pembuatan event konser baru, pembuatan seksi seat map, penentuan harga tiket & kuota.", "Selesai", "FrontEnd Dev"),
    (9, "Organizer AI Business Intelligence Panel", "Panel pembuat deskripsi event otomatis berbasis AI dan grafik proyeksi tren penjualan tiket.", "Selesai", "FrontEnd Dev"),
    (10, "Scanner Mobile App / E-Ticket Check-in", "Antarmuka kamera pemindai QR Code E-Ticket untuk petugas pintu masuk acara konser.", "Selesai", "FrontEnd Dev"),
    (11, "Super Admin Panel & Verification Dashboard", "Halaman verifikasi legalitas organizer, pengelolaan master venue, dan pemantauan transaksi platform.", "Selesai", "FrontEnd Dev"),
    (12, "Super Admin AI Analytics Dashboard", "Panel analitik performa platform berbasis AI dan grafik deteksi transaksi yang mencurigakan.", "Selesai", "FrontEnd Dev")
]

row_idx = 4
for no, mod, desc, status, pic in fe_tasks:
    ws5.cell(row=row_idx, column=1, value=no).alignment = align_center
    ws5.cell(row=row_idx, column=2, value=mod).font = font_bold
    ws5.cell(row=row_idx, column=3, value=desc).font = font_regular
    
    cell_status = ws5.cell(row=row_idx, column=4, value=status)
    cell_status.alignment = align_center
    cell_status.font = font_bold
    cell_status.fill = fill_green_status if status == "Selesai" else fill_blue_status
    
    ws5.cell(row=row_idx, column=5, value=pic).alignment = align_center
    
    ws5.cell(row=row_idx, column=2).alignment = align_left
    ws5.cell(row=row_idx, column=3).alignment = align_left
    
    for c in range(1, 6):
        ws5.cell(row=row_idx, column=c).border = thin_border
    ws5.row_dimensions[row_idx].height = 32
    row_idx += 1


# Auto-fit columns for all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignore title rows when calculating width
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or '')
            if '\n' in val_str:
                lines = val_str.split('\n')
                line_len = max(len(l) for l in lines)
                if line_len > max_len: max_len = line_len
            else:
                if len(val_str) > max_len: max_len = len(val_str)
        # Cap column width
        adjusted_width = min(max(max_len + 4, 12), 80)
        sheet.column_dimensions[col_letter].width = adjusted_width

# Save new workbook
output_path = r'C:\KetTiket\Deskripsi_dan_List_Tasks_KetTiket.xlsx'
wb.save(output_path)
print("Workbook successfully created at:", output_path)
